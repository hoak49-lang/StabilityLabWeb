"""Strict tabular import with explicit column mapping and source row IDs."""
import bootstrap
import csv
import io
import hashlib
import zipfile
from pathlib import Path
from openpyxl import load_workbook, Workbook
from engine import AnalysisError, number

MAX_BYTES = 8*1024*1024


def parse_file(data, filename, sheet=None):
    if len(data)>MAX_BYTES:
        raise AnalysisError('Tệp vượt 8 MB. Hãy tách riêng bảng độ ổn định cần phân tích.')
    extension=Path(filename).suffix.lower()
    sheets=[]
    if extension=='.csv':
        try:
            content=data.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise AnalysisError('CSV cần mã hóa UTF-8. Hãy lưu dạng CSV UTF-8 trong Excel.')
        try:
            dialect=csv.Sniffer().sniff(content[:8192],delimiters=',;\t')
        except csv.Error:
            dialect=csv.excel
        matrix=list(csv.reader(io.StringIO(content),dialect))
    elif extension=='.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist())>80*1024*1024 or len(z.infolist())>3000:
                    raise AnalysisError('Workbook giải nén quá lớn.')
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
            sheets=wb.sheetnames
            if sheet is not None and sheet not in sheets:
                raise AnalysisError('Tên trang tính không tồn tại.')
            ws=wb[sheet or sheets[0]]
            if ws.max_row>10001 or ws.max_column>80:
                raise AnalysisError('Bảng nhập tối đa 10.000 dòng và 80 cột. Xóa định dạng dư ngoài bảng nếu cần.')
            matrix=[]
            for row in ws.iter_rows():
                values=[]
                for cell in row:
                    if cell.data_type=='f':
                        raise AnalysisError(f'Ô {cell.coordinate} chứa công thức. Xuất một bản kết quả đã kiểm tra dưới dạng giá trị để nhập.')
                    if cell.data_type=='e':
                        raise AnalysisError(f'Ô {cell.coordinate} có lỗi Excel.')
                    value=cell.value
                    if value is not None and not isinstance(value,(str,int,float,bool)):
                        value=str(value)
                    values.append(value)
                matrix.append(values)
            wb.close()
        except AnalysisError:
            raise
        except Exception:
            raise AnalysisError('Không đọc được XLSX. Dùng workbook không mã hóa, không macro, và kiểm tra tệp gốc.')
    else:
        raise AnalysisError('Chỉ nhận .csv và .xlsx. Hãy chuyển .xls sang .xlsx.')
    while matrix and all(v is None or str(v).strip()=='' for v in matrix[-1]):
        matrix.pop()
    if len(matrix)<2 or len(matrix)>10001:
        raise AnalysisError('Cần dòng tiêu đề và từ 1 đến 10.000 dòng dữ liệu.')
    headers=[str(v).strip() if v is not None else '' for v in matrix[0]]
    while headers and headers[-1]=='':
        headers.pop()
    if not headers or any(not h for h in headers) or len(set(headers))!=len(headers):
        raise AnalysisError('Dòng đầu phải chứa tên cột duy nhất, không để trống giữa các cột.')
    if len(headers)>80 or any(len(h)>160 for h in headers):
        raise AnalysisError('Tên cột hoặc số cột vượt giới hạn.')
    rows=[]
    blank_rows=[]
    for index,row in enumerate(matrix[1:],2):
        if all(v is None or str(v).strip()=='' for v in row):
            blank_rows.append(index)
            continue
        if len(row)>len(headers) and any(v is not None and str(v).strip() for v in row[len(headers):]):
            raise AnalysisError(f'Dòng {index} có dữ liệu ngoài cột tiêu đề.')
        row=(row+[None]*len(headers))[:len(headers)]
        if any(isinstance(v,str) and len(v)>5000 for v in row):
            raise AnalysisError(f'Dòng {index}: nội dung ô quá dài.')
        rows.append(dict(row_id=index,cells=row))
    return dict(filename=Path(filename).name,sha256=hashlib.sha256(data).hexdigest(),
                sheets=sheets,sheet=sheet or (sheets[0] if sheets else None),headers=headers,
                rows=rows,blank_rows=blank_rows)


def map_rows(table, mapping, decimal='dot'):
    if decimal not in ('dot','comma'):
        raise AnalysisError('Dấu thập phân không hợp lệ.')
    idx={}
    for field in ('batch','time','value','attribute','condition','strength','pack'):
        name=mapping.get(field)
        if not name and field in ('strength','pack'):
            aliases={'strength':['strength','hàm lượng','ham luong'],'pack':['pack','package','quy cách','quy cach']}
            matches=[h for h in table['headers'] if h.lower() in aliases[field]]
            if matches:name=matches[0]
        if field in ('batch','time','value') and not name:
            raise AnalysisError('Hãy chọn cột lô, thời gian và kết quả.')
        if name:
            if name not in table['headers']:
                raise AnalysisError('Cột đã chọn không tồn tại.')
            idx[field]=table['headers'].index(name)
    if len(set(idx.values()))!=len(idx):
        raise AnalysisError('Mỗi vai trò phải dùng một cột riêng.')
    result=[]
    for row in table['rows']:
        record=dict(row_id=row['row_id'],attribute='Chỉ tiêu',condition='Dài hạn')
        for field,i in idx.items():
            v=row['cells'][i]
            if field in ('time','value'):
                if isinstance(v,str) and decimal=='comma':
                    if '.' in v and ',' in v:
                        raise AnalysisError(f'Dòng {row["row_id"]}: không hỗ trợ dấu phân cách hàng nghìn. Chuẩn hóa dữ liệu trước khi nhập.')
                    v=v.replace(',','.')
                record[field]=number(v,f'Dòng {row["row_id"]}, {field}')
            else:
                if v is None or not str(v).strip():
                    raise AnalysisError(f'Dòng {row["row_id"]}: {field} bị trống. Không tự xóa dòng.')
                record[field]=str(v).strip()
        record['study_group']=__import__('json').dumps([record.get('strength',''),record.get('pack','')],ensure_ascii=False)
        result.append(record)
    return result


def safe_cell(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):
        return "'"+value
    return value


def workbook_bytes(sheets):
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook()
    wb.remove(wb.active)
    for title,rows in sheets:
        ws=wb.create_sheet(title[:31])
        for row in rows:
            ws.append([safe_cell(v) for v in row])
        ws.freeze_panes='A2'
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='173C46')
        for col in ws.columns:
            letter=col[0].column_letter
            ws.column_dimensions[letter].width=min(65,max(16,max(len(str(c.value or '')) for c in col[:150])+2))
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,float):
                    cell.number_format='0.000000'
    out=io.BytesIO()
    wb.save(out)
    return out.getvalue()
