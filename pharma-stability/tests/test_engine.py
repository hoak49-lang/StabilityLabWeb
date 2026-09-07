import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import bootstrap
import unittest
import json,zipfile,io
import numpy as np
from scipy import stats
from engine import analyze,AnalysisError,extrapolation_cap
from dataio import parse_file,map_rows,workbook_bytes


def generated(slopes=(-.2,-.2,-.2), intercepts=(100,100,100), noise=.05):
    rng=np.random.default_rng(813)
    times=np.array([0,3,6,9,12,18,24.])
    basis=np.column_stack([np.ones(7),times,times**2])
    rows=[]
    for i,(a,b) in enumerate(zip(intercepts,slopes)):
        raw=rng.normal(size=7)
        error=raw-basis@np.linalg.lstsq(basis,raw,rcond=None)[0]
        error=error/np.linalg.norm(error)*noise*np.sqrt(7)
        rows.extend(dict(batch=str(i+1),time=float(t),value=float(a+b*t+e)) for t,e in zip(times,error))
    return rows


class StatisticsTests(unittest.TestCase):
    def test_minitab_published_fixed_batch_example(self):
        z=zipfile.ZipFile(Path(__file__).parent/'ShelfLife.MWX')
        c=json.loads(z.read('/sheets/0/sheet.json'))['Data']['Columns']
        cols={x['WorksheetVarBody']['Name']:x['WorksheetVarBody']['VarData']['VarDataBody'] for x in c}
        rows=[dict(batch=b,time=t,value=v) for b,t,v in zip(cols['Batch']['TextData'],cols['Month']['NumericData'],cols['Drug%']['NumericData']) if v<1e29]
        self.assertEqual(len(rows),40)  # 5 documented Minitab missing-value sentinels
        r=analyze(rows,dict(lower=90,direction='decrease',horizon=120))
        self.assertEqual(r['model'],'separate')
        self.assertEqual(r['df_error'],30)
        self.assertAlmostEqual(r['rmse'],.594983,places=6)
        self.assertAlmostEqual(r['slope_test']['p'],.048,delta=.0005)
        self.assertAlmostEqual(r['statistical_months'],54.790,delta=.001)
        for actual,expected in zip(r['coefficients'],[83.552,54.790,57.492,60.898,66.854]):
            self.assertAlmostEqual(actual['crossing'],expected,delta=.001)

    def test_pooling_paths(self):
        self.assertEqual(analyze(generated(),dict(lower=90))['model'],'pooled')
        self.assertEqual(analyze(generated(intercepts=(100,102,104)),dict(lower=90))['model'],'common_slope')
        self.assertEqual(analyze(generated(slopes=(-.1,-.3,-.5)),dict(lower=90))['model'],'separate')

    def test_single_batch_matches_closed_form(self):
        rows=generated()[:7]
        r=analyze(rows,dict(lower=95,horizon=100))
        x=np.array([p['time'] for p in rows]); y=np.array([p['value'] for p in rows])
        regression=stats.linregress(x,y)
        mse=sum((y-regression.intercept-regression.slope*x)**2)/(len(x)-2)
        for point in r['series'][0]['points'][::31]:
            t=point['time']; se=np.sqrt(mse*(1/len(x)+(t-x.mean())**2/sum((x-x.mean())**2)))
            expected=regression.intercept+regression.slope*t-stats.t.ppf(.95,len(x)-2)*se
            self.assertAlmostEqual(point['lower'],expected,places=10)

    def test_both_sides_is_more_conservative(self):
        rows=generated()
        one=analyze(rows,dict(lower=95,upper=105,direction='decrease'))
        both=analyze(rows,dict(lower=95,upper=105,direction='both'))
        self.assertLess(both['statistical_months'],one['statistical_months'])
        self.assertEqual(both['confidence_sides'],2)

    def test_increasing_and_log(self):
        rows=generated(slopes=(.2,.2,.2),intercepts=(1,1,1))
        r=analyze(rows,dict(upper=6,direction='increase'))
        self.assertTrue(20<r['statistical_months']<25)
        logrows=[dict(r,value=np.exp(r['value']/100)) for r in generated()]
        r=analyze(logrows,dict(lower=2.5,transform='log'))
        self.assertTrue(30<r['statistical_months']<50)

    def test_zero_and_no_crossing(self):
        r=analyze(generated(),dict(lower=101))
        self.assertEqual(r['statistical_months'],0)
        r=analyze(generated(),dict(lower=1,horizon=120))
        self.assertIsNone(r['statistical_months'])
        self.assertEqual(r['statistical_status'],'beyond_search')
        self.assertEqual(r['conditional_limit_months'],24)

    def test_invalid_data(self):
        for value in [None,'<0.1',float('nan'),True]:
            rows=generated(); rows[0]['value']=value
            with self.assertRaises(AnalysisError): analyze(rows,dict(lower=90))
        rows=generated(); rows[0]['condition']='Accelerated'
        with self.assertRaises(AnalysisError): analyze(rows,dict(lower=90))
        with self.assertRaises(AnalysisError): analyze(generated(),dict(lower=110,upper=90))
        with self.assertRaises(AnalysisError): analyze(generated(),dict(lower=0,transform='log'))

    def test_no_silent_outlier_removal(self):
        rows=generated(); rows[4]['value']=50
        r=analyze(rows,dict(lower=90))
        self.assertEqual(r['n'],len(rows))
        self.assertIsNone(r['candidate_months'])
        self.assertTrue(any('ngoài giới hạn' in x for x in r['blockers']))

    def test_no_auto_proposal(self):
        r=analyze(generated(),dict(lower=90))
        self.assertIsNone(r['candidate_months'])
        self.assertEqual(r['status'],'review_required')

    def test_candidate_and_conservative_rounding(self):
        review=dict(long_term_confirmed=True,design_confirmed=True,model_reviewed=True,
                    other_attributes_reviewed=True,reviewer='Test analyst',justification='Synthetic validation test only.',product='Test')
        r=analyze(generated(),dict(lower=95),review)
        self.assertEqual(r['blockers'],[])
        self.assertEqual(r['candidate_months'],24)

    def test_degenerate_model_does_not_claim_confidence(self):
        rows=[dict(r,value=100-.2*r['time']) for r in generated()]
        r=analyze(rows,dict(lower=90))
        self.assertIsNone(r['candidate_months'])
        self.assertTrue(any('Sai số dư' in b for b in r['blockers']))

    def test_curvature_block(self):
        rows=[dict(r,value=r['value']-.02*r['time']**2) for r in generated()]
        r=analyze(rows,dict(lower=70))
        self.assertTrue(any('dấu hiệu cong' in b for b in r['blockers']))

    def test_unbalanced_duration(self):
        rows=[r for r in generated() if r['batch']!='1' or r['time']<=12]
        self.assertEqual(analyze(rows,dict(lower=90))['observed_common_months'],12)

    def test_data_order_invariance(self):
        rows=generated(slopes=(-.1,-.3,-.5))
        a=analyze(rows,dict(lower=90)); b=analyze(list(reversed(rows)),dict(lower=90))
        self.assertAlmostEqual(a['statistical_months'],b['statistical_months'],places=8)


class ExtrapolationTests(unittest.TestCase):
    def test_decision_branches(self):
        review=dict(allow_extrapolation=True,storage='room',accelerated='no_change',accelerated_months=6,
                    supporting_data=True,evidence='Reviewed development batch stability report.')
        self.assertEqual(extrapolation_cap(12,review)['months'],24)
        self.assertEqual(extrapolation_cap(24,review)['months'],36)
        self.assertEqual(extrapolation_cap(12,dict(review,storage='refrigerated'))['months'],18)
        self.assertEqual(extrapolation_cap(12,dict(review,storage='frozen'))['months'],12)
        self.assertEqual(extrapolation_cap(12,dict(review,accelerated='unknown'))['months'],12)
        self.assertEqual(extrapolation_cap(12,dict(review,supporting_data=False))['months'],12)
        self.assertEqual(extrapolation_cap(12,dict(review,accelerated_months=3))['months'],12)
        changed=dict(review,accelerated='change',intermediate='no_change',intermediate_months=6)
        self.assertEqual(extrapolation_cap(12,changed)['months'],18)
        self.assertEqual(extrapolation_cap(12,dict(changed,storage='refrigerated'))['months'],12)
        self.assertEqual(extrapolation_cap(12,dict(changed,intermediate='change'))['months'],12)


class ImportTests(unittest.TestCase):
    def test_csv_comma_decimal(self):
        t=parse_file('Lô;Tháng;Kết quả\nA;0;100,2\nA;3;99,5\n'.encode(),'test.csv')
        r=map_rows(t,dict(batch='Lô',time='Tháng',value='Kết quả'),'comma')
        self.assertEqual(r[0]['value'],100.2)
    def test_missing_value_not_dropped(self):
        t=parse_file(b'b,t,v\nA,0,100\nA,3,\n','test.csv')
        with self.assertRaises(AnalysisError): map_rows(t,dict(batch='b',time='t',value='v'))
    def test_blank_rows_explicit(self):
        t=parse_file(b'b,t,v\nA,0,100\n,,\nA,3,99\n','test.csv')
        self.assertEqual(t['blank_rows'],[3])
    def test_formula_rejected_and_export_safe(self):
        from openpyxl import Workbook,load_workbook
        wb=Workbook(); ws=wb.active; ws.append(['batch','time','value']); ws.append(['A',0,'=100+1'])
        out=io.BytesIO(); wb.save(out)
        with self.assertRaises(AnalysisError): parse_file(out.getvalue(),'x.xlsx')
        result=workbook_bytes([('test',[['label'],['=HYPERLINK("bad")']])])
        exported=load_workbook(io.BytesIO(result))
        self.assertNotEqual(exported.active['A2'].data_type,'f')
    def test_duplicate_mapping(self):
        t=parse_file(b'b,t,v\nA,0,100\n','x.csv')
        with self.assertRaises(AnalysisError): map_rows(t,dict(batch='b',time='t',value='t'))


if __name__=='__main__': unittest.main(verbosity=2)
